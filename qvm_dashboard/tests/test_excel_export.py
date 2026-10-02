"""Tests for the Via-to-Pad Excel export (chart embedding + layout)."""

from io import BytesIO
import os

import numpy as np
import pandas as pd
import pytest

pytest.importorskip('kaleido')  # chart embedding renders PNGs via kaleido


def test_ensure_chrome_on_path_sets_common_windows_locations(monkeypatch):
    import src.excel_export as excel_export

    def fake_isfile(path):
        return path == r'C:\Program Files\Google\Chrome\Application\chrome.exe'

    monkeypatch.setattr(excel_export.os.path, 'isfile', fake_isfile)
    monkeypatch.delenv('PATH', raising=False)
    monkeypatch.delenv('GOOGLE_CHROME_BIN', raising=False)
    monkeypatch.delenv('CHROME_BIN', raising=False)

    excel_export._ensure_chrome_on_path()

    assert os.environ['GOOGLE_CHROME_BIN'] == r'C:\Program Files\Google\Chrome\Application\chrome.exe'
    assert os.environ['CHROME_BIN'] == r'C:\Program Files\Google\Chrome\Application\chrome.exe'
    assert r'C:\Program Files\Google\Chrome\Application' in os.environ.get('PATH', '')


@pytest.fixture(autouse=True)
def _silence_streamlit(monkeypatch):
    """The exporter calls st.sidebar.warning on embed failure; stub it."""
    import streamlit as st
    monkeypatch.setattr(st.sidebar, 'warning', lambda *a, **k: None)


def _vtp_frame(sides, settings):
    cn = settings['COLUMN_NAMES']
    rng = np.random.default_rng(1)
    grid = (list(range(11, 15)) + list(range(21, 25)) +
            list(range(31, 35)) + list(range(41, 45)))
    frames = []
    for side in sides:
        n = 16
        bu = 'BU-01' if side == 'F' else 'BU-02'
        frames.append(pd.DataFrame({
            cn['grid_id']: grid,
            cn['location']: [f'p{i}' for i in range(n)],
            cn['inner_diameter']: np.full(n, 0.02),
            cn['outer_diameter']: np.full(n, 0.3125),
            cn['coord_x']: rng.uniform(50, 450, n),
            cn['coord_y']: rng.uniform(50, 450, n),
            cn['x_distance']: rng.normal(0, 0.004, n),
            cn['y_distance']: rng.normal(0, 0.004, n),
            cn['ptv_distance']: np.abs(rng.normal(0.005, 0.002, n)),
            'Build Up': [bu] * n,
            'Side': [side] * n, 'Panel': ['20'] * n,
            'Process': ['Via to Pad'] * n,
        }))
    return pd.concat(frames, ignore_index=True)


def _export(settings, sides, meta=None):
    """Run the real export entry point over a Via-to-Pad frame."""
    import openpyxl
    from src.excel_export import write_unified_excel
    df = _vtp_frame(sides, settings)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(w, df.copy(), settings,
                            meta or {'panel': '20', 'process': 'Via to Pad'})
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def _analytics_sheets(wb):
    return [n for n in wb.sheetnames if n.startswith('Analytics')]


def test_each_build_up_gets_analytics_sheet_with_four_charts(settings):
    # _vtp_frame tags side F as BU-01 and side B as BU-02.
    wb = _export(settings, ['F', 'B'])
    sheets = _analytics_sheets(wb)
    assert sorted(sheets) == ['Analytics VTP BU-01', 'Analytics VTP BU-02']
    for name in sheets:
        assert len(wb[name]._images) == 4, f'{name} should embed all 4 charts'


def test_single_build_up_gets_one_analytics_sheet(settings):
    wb = _export(settings, ['F'])
    assert _analytics_sheets(wb) == ['Analytics VTP BU-01']
    images = wb['Analytics VTP BU-01']._images
    assert len(images) == 4
    assert {im.anchor._from.row for im in images} == {2}
    assert len({im.anchor._from.col for im in images}) == 4


def test_same_build_up_front_back_embed_eight_side_specific_charts(settings):
    import openpyxl
    from src.excel_export import write_unified_excel

    df = _vtp_frame(['F', 'B'], settings)
    df['Build Up'] = 'BU-01'

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(w, df.copy(), settings,
                            {'panel': '20', 'process': 'Via to Pad'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    assert _analytics_sheets(wb) == ['Analytics VTP BU-01']
    images = wb['Analytics VTP BU-01']._images
    assert len(images) == 8

    # Two side rows, each with four charts across the sheet.
    cols = sorted({im.anchor._from.col for im in images})
    rows = sorted({im.anchor._from.row for im in images})
    assert len(cols) == 4
    assert len(rows) == 2
    report = wb['Analytics VTP BU-01']
    assert report['A2'].value == 'Side: Front'
    assert report['A74'].value == 'Side: Back'


def test_measurement_sheet_has_no_duplicated_charts(settings):
    """Regression: bullseye + bars moved to Analytics; data sheet keeps only
    the Panel Map (≤1 image), previously 3."""
    wb = _export(settings, ['F', 'B'])
    for name in wb.sheetnames:
        if name.startswith('Via to Pad'):
            assert len(wb[name]._images) <= 1, f'{name} must not duplicate charts'


def test_analytics_charts_do_not_overlap(settings):
    """Regression: the four-chart row must be spaced so images do not collide."""
    wb = _export(settings, ['F'])
    anchors = [(im.anchor._from.col, im.anchor._from.row)
               for im in wb['Analytics VTP BU-01']._images]
    cols = sorted({c for c, _ in anchors})
    rows = sorted({r for _, r in anchors})
    assert len(cols) == 4 and len(rows) == 1, 'charts should form one row'
    # neighboring charts must remain separated by a safe gap
    assert cols[1] - cols[0] >= 13, 'second chart too close (overlap risk)'
    assert cols[2] - cols[1] >= 13, 'third chart too close (overlap risk)'
    assert cols[3] - cols[2] >= 13, 'fourth chart too close (overlap risk)'


def test_analytics_sheet_falls_back_when_chart_render_fails(settings, monkeypatch):
    """A failed render must still leave a placeholder image, not an empty sheet."""
    import openpyxl
    from src import excel_export

    df = _vtp_frame(['F'], settings)
    # Patch where it is looked up (qvm sheet builder), not where it is defined.
    monkeypatch.setattr(
        excel_export.qvm, '_fig_to_xl_image',
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError('kaleido blocked')))

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        excel_export.write_unified_excel(w, df.copy(), settings,
                                         {'panel': '20', 'process': 'Via to Pad'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    assert 'Analytics VTP BU-01' in wb.sheetnames
    assert len(wb['Analytics VTP BU-01']._images) == 4


def _report_rows(wb):
    """Report sheet as a list of {header: value} dicts."""
    report = wb['Report']
    headers = [c.value for c in report[1]]
    return [dict(zip(headers, [c.value for c in report[r]]))
            for r in range(2, report.max_row + 1)]


def test_report_sheet_preserves_row_level_build_up(settings):
    wb = _export(settings, ['F', 'B'])
    bu_values = {row['Buildup'] for row in _report_rows(wb)}
    assert {'BU-01', 'BU-02'}.issubset(bu_values)


def test_report_row_provenance_stays_with_its_own_fiducial(settings):
    """Regression: Side and Buildup were taken positionally from the parsed
    frame (grouped by file) and written onto the display frame (sorted by Grid
    ID). With both sides in scope the two orders interleave, and half the rows
    were labelled with the other side's values.

    _vtp_frame ties side F to BU-01 and side B to BU-02, so any row pairing
    F with BU-02 (or B with BU-01) is a misattribution.
    """
    wb = _export(settings, ['F', 'B'])
    rows = _report_rows(wb)

    expected = {'F': 'BU-01', 'B': 'BU-02'}
    mismatched = [r for r in rows if expected[r['Side']] != r['Buildup']]
    assert not mismatched, f'{len(mismatched)} of {len(rows)} rows misattributed'

    # Both sides must actually be represented, or the check above is vacuous.
    assert {r['Side'] for r in rows} == {'F', 'B'}


def test_report_panel_comes_from_the_data(settings):
    """Panel was hardcoded to the (now always empty) job-metadata value even
    though every row carries its own."""
    wb = _export(settings, ['F'])
    assert {r['Panel'] for r in _report_rows(wb)} == {'20'}


def _mixed_process_frame(settings):
    vtp = _vtp_frame(['F'], settings)
    vtp['Process'] = 'Via to Pad'
    ptv = _vtp_frame(['B'], settings)
    ptv['Process'] = 'Pad to Via'
    return pd.concat([vtp, ptv], ignore_index=True)


def _valid_loc_frame(settings, process: str, build_up: str, side: str = 'F'):
    nominal = settings['CORE_SHRINKAGE']['NOMINAL_COORDS']
    cn = settings['COLUMN_NAMES']
    grid_map = settings['GRID_MAPPING']
    rows = []
    for loc, xy in nominal.items():
        rows.append({
            cn['grid_id']: int(grid_map[loc]),
            cn['location']: loc,
            cn['inner_diameter']: 0.02,
            cn['outer_diameter']: 0.3125,
            cn['coord_x']: float(xy['x']),
            cn['coord_y']: float(xy['y']),
            cn['x_distance']: 0.001,
            cn['y_distance']: 0.001,
            cn['ptv_distance']: 0.002,
            'Build Up': build_up,
            'Side': side,
            'Panel': '20',
            'Process': process,
        })
    return pd.DataFrame(rows)


def test_unified_export_writes_single_workbook_with_mixed_process_data(settings):
    import openpyxl
    from src.excel_export import write_unified_excel

    df = _mixed_process_frame(settings)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(
            w,
            df.copy(),
            settings,
            {'panel': '20', 'process': 'MIXED', 'build_up': 'BU-ALL'},
            warnings=['sample warning'],
        )
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    assert 'Via to Pad' in wb.sheetnames
    assert 'Pad to Via' in wb.sheetnames
    assert 'Via to Pad - F' not in wb.sheetnames
    assert 'Pad to Via - B' not in wb.sheetnames
    assert 'Report' in wb.sheetnames
    assert any(name.startswith('Warnings') for name in wb.sheetnames)


def test_unified_export_process_sheets_embed_panel_map_when_available(settings):
    import openpyxl
    from src import excel_export

    export_cfg = settings.get('EXCEL_EXPORT', {})
    rel_path = export_cfg.get('panel_map_image', 'assets/Panel_Map.png')
    abs_path = os.path.join(excel_export.PACKAGE_ROOT, rel_path)

    if not os.path.isfile(abs_path):
        pytest.skip('Panel map image not available in test environment')

    df = _mixed_process_frame(settings)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        excel_export.write_unified_excel(w, df.copy(), settings, {'panel': '20'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    assert len(wb['Via to Pad']._images) >= 1
    assert len(wb['Pad to Via']._images) >= 1


def test_unified_export_panel_map_is_anchored_right_of_data(settings):
    import openpyxl
    from src import excel_export

    export_cfg = settings.get('EXCEL_EXPORT', {})
    rel_path = export_cfg.get('panel_map_image', 'assets/Panel_Map.png')
    abs_path = os.path.join(excel_export.PACKAGE_ROOT, rel_path)

    if not os.path.isfile(abs_path):
        pytest.skip('Panel map image not available in test environment')

    df = _mixed_process_frame(settings)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        excel_export.write_unified_excel(w, df.copy(), settings, {'panel': '20'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    ws = wb['Via to Pad']
    assert ws._images, 'Expected panel-map image on Via to Pad sheet'
    image_anchor_col = ws._images[0].anchor._from.col
    # openpyxl columns are 0-based in anchors; image should be at least one column
    # to the right of the final data column.
    assert image_anchor_col >= ws.max_column


def test_unified_export_creates_analytics_per_process_and_build_up(settings):
    import openpyxl
    from src.excel_export import write_unified_excel

    df = _mixed_process_frame(settings)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(w, df.copy(), settings, {'panel': '20'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    analytics_sheets = [name for name in wb.sheetnames if name.startswith('Analytics')]
    assert any('VTP' in name and 'BU-01' in name for name in analytics_sheets)
    assert any('PTV' in name and 'BU-02' in name for name in analytics_sheets)
    for name in analytics_sheets:
        assert len(wb[name]._images) == 4


def test_unified_export_report_has_process_specific_usl_columns(settings):
    import openpyxl
    from src.excel_export import write_unified_excel

    df = _mixed_process_frame(settings)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(w, df.copy(), settings, {'panel': '20'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    headers = [cell.value for cell in wb['Report'][1]]
    assert 'VtP_USL' in headers
    assert 'PtV_USL' in headers


def test_unified_export_process_sheet_includes_annular_ring(settings):
    import openpyxl
    from src.excel_export import write_unified_excel

    df = _mixed_process_frame(settings)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(w, df.copy(), settings, {'panel': '20'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    headers = [cell.value for cell in wb['Via to Pad'][1]]
    assert 'Annular Ring' in headers


def test_unified_export_handles_malformed_grid_ids(settings):
    import openpyxl
    from src.excel_export import write_unified_excel

    df = _mixed_process_frame(settings)
    grid_col = settings['COLUMN_NAMES']['grid_id']
    # Object dtype first: pandas >= 2 refuses string setitem into an int64
    # column, but real uploads arrive with mixed-type Grid IDs.
    df[grid_col] = df[grid_col].astype(object)
    df.loc[df.index[:4], grid_col] = ['A', 'B2', '3X', '']

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(w, df.copy(), settings, {'panel': '20'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    assert 'Report' in wb.sheetnames
    analytics_sheets = [name for name in wb.sheetnames if name.startswith('Analytics')]
    assert analytics_sheets


def test_unified_export_writes_combined_shrinkage_summary_sheet(settings):
    import openpyxl
    from src.excel_export import write_unified_excel

    df = pd.concat([
        _valid_loc_frame(settings, process='Via to Pad', build_up='BU-01', side='F'),
        _valid_loc_frame(settings, process='Pad to Via', build_up='Unknown', side='B'),
    ], ignore_index=True)

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(w, df.copy(), settings, {'panel': '20'})
    buf.seek(0)
    wb = openpyxl.load_workbook(buf)

    shrink_sheets = [name for name in wb.sheetnames if name.startswith('Shrinkage')]
    assert len(shrink_sheets) == 1
    ws = wb[shrink_sheets[0]]

    values = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=1):
        cell_val = row[0].value
        if isinstance(cell_val, str):
            values.append(cell_val)

    assert any('VTP | Build Up: BU-01' in v for v in values)
    assert any('PTV | Build Up: Unknown' in v for v in values)


def test_shrinkage_sheet_uses_registration_design_and_company_numbers(settings):
    import openpyxl
    from src.excel_export import write_unified_excel
    from src.registration_shrinkage import design_for

    cn = settings['COLUMN_NAMES']

    def frame(build_up, process, scale):
        design = design_for(settings, 'BU-01', process)
        out = _valid_loc_frame(settings, process=process, build_up=build_up)
        out[cn['coord_x']] = [design[loc]['x'] * scale for loc in out[cn['location']]]
        out[cn['coord_y']] = [design[loc]['y'] * scale for loc in out[cn['location']]]
        return out

    df = pd.concat([frame('BU-01', 'Via to Pad', 0.9999),
                    frame('BU-02', 'Pad to Via', 1.0)], ignore_index=True)
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_unified_excel(w, df.copy(), settings, {'panel': '20'})
    buf.seek(0)
    ws = openpyxl.load_workbook(buf)['Shrinkage Summary']
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    text = [str(c) for row in rows for c in row if c is not None]

    assert 'Method: Company (distance).' in rows[1][0]
    header = next(row for row in rows if row and row[0] == 'Process')
    values = rows[rows.index(header) + 1]
    summary = dict(zip(header, values))
    assert summary['Process'] == 'VTP' and summary['Build Up'] == 'BU-01'
    assert summary['FP X (ppm)'] == pytest.approx(-100.0, abs=0.01)
    assert summary['FP X (um)'] == pytest.approx(-23.25, abs=0.01)
    assert summary['Alert Count'] == 0
    assert not any('%' in t for t in header if isinstance(t, str))

    # BU-02 Pad to Via has no entry: the equivalent BU-01 design stands in, and says so.
    assert any(t.startswith('PTV | Build Up: BU-02') for t in text)
    assert any('No design positions configured for BU-02 / Pad to Via; those of BU-01 / '
               'Pad to Via are used' in t for t in text)
    second = dict(zip(header, rows[rows.index(header) + 2]))
    assert second['Build Up'] == 'BU-02' and second['FP X (ppm)'] == pytest.approx(0.0, abs=0.01)
    assert any(t == 'Via to Pad, via centres, Company (distance)' for t in text)
