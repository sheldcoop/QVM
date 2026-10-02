"""Tests for the Core Shrinkage parser, calculations, frame builders, and Excel.

Fixtures generate "measured" data by applying a KNOWN affine transform (scale /
rotation / datum translation), optionally with an injected outlier dot — NOT by
copying nominal verbatim. That way the tests exercise the real recovery path and
guard the −471 mm numbering-swap failure mode, instead of being self-fulfilling.
"""

import math
from io import BytesIO

import numpy as np
import pandas as pd
import pytest

from src.core_shrinkage_parser import (
    parse_cs_filename,
    parse_core_shrinkage_content,
    sort_by_grid,
    layer_frame,
    both_frame,
    compute_qp_shrinkage,
    compute_scale_factor,
    compute_affine_analysis,
    compute_shrinkage_summary,
    summary_frame,
    qp_shrinkage_um,
    affine_residuals_um,
    flag_outliers,
    compute_front_back_registration,
    compute_deviation_field,
    compute_centered_deviation_table,
    analyze_layer,
    _rigid_fit,
    _classify_vector,
)

# Nominal design coordinates from settings.yaml CORE_SHRINKAGE.NOMINAL_COORDS
_NOMINAL = {
    'UL_1': {'x': -222.700, 'y':  239.875},
    'UL_2': {'x': -222.700, 'y':    4.125},
    'UL_3': {'x':  -22.050, 'y':    4.125},
    'UL_4': {'x':  -22.050, 'y':  239.875},
    'LL_1': {'x': -222.700, 'y':   -4.125},
    'LL_2': {'x': -222.700, 'y': -239.875},
    'LL_3': {'x':  -22.050, 'y': -239.875},
    'LL_4': {'x':  -22.050, 'y':   -4.125},
    'LR_1': {'x':   20.300, 'y':   -4.125},
    'LR_2': {'x':   20.300, 'y': -239.875},
    'LR_3': {'x':  220.950, 'y': -239.875},
    'LR_4': {'x':  220.950, 'y':   -4.125},
    'UR_1': {'x':   20.300, 'y':  239.875},
    'UR_2': {'x':   20.300, 'y':    4.125},
    'UR_3': {'x':  220.950, 'y':    4.125},
    'UR_4': {'x':  220.950, 'y':  239.875},
}

_QUADS = ('UL', 'LL', 'LR', 'UR')
# UL=1x, LL=2x, LR=3x, UR=4x  →  matches GRID_MAPPING in settings.yaml
GRID = {f'{q}_{n}': (i + 1) * 10 + n
        for i, q in enumerate(_QUADS) for n in range(1, 5)}


def _measured_df(layer: str, *, sx: float = 1.0, sy: float = 1.0,
                 theta_deg: float = 0.0, tx: float = 0.0, ty: float = 0.0,
                 outlier: str = None, outlier_dxy=(0.0, 0.0)) -> pd.DataFrame:
    """Build measured data as transform(nominal): scale → rotate → translate,
    with an optional single outlier dot displaced by outlier_dxy (mm)."""
    t = math.radians(theta_deg)
    cos_t, sin_t = math.cos(t), math.sin(t)
    rows = []
    for loc, c in _NOMINAL.items():
        quad, num = loc.split('_')
        xs, ys = c['x'] * sx, c['y'] * sy
        mx = cos_t * xs - sin_t * ys + tx
        my = sin_t * xs + cos_t * ys + ty
        if loc == outlier:
            mx += outlier_dxy[0]
            my += outlier_dxy[1]
        rows.append({
            'location': loc, 'quadrant': quad, 'point_num': int(num),
            'coord_x': mx, 'coord_y': my, 'diameter': 0.370, 'pts_count': 150,
            'layer': layer, 'Process': 'Core Shrinkage', 'grid_id': GRID.get(loc),
        })
    return pd.DataFrame(rows)


def _nominal_df(layer: str) -> pd.DataFrame:
    """Measured == nominal (identity transform)."""
    return _measured_df(layer)


def _content(prefix: str) -> str:
    """Build a 16-circle Core Shrinkage file. prefix='Pad_' (Front) or '' (Back)."""
    lines, idx = [], 1
    for q in _QUADS:
        for n in range(1, 5):
            lines.append(f"Circle: {prefix}{q}_{n}(ID:{idx}, From 150 Pts.)")
            lines.append(f"     Coord. X =        {10.0 + idx:.6f}")
            lines.append(f"     Coord. Y =        {20.0 + idx:.6f}")
            lines.append(f"     Diameter =         0.370000")
            idx += 1
    return "\n".join(lines)


def _parsed(layer: str) -> pd.DataFrame:
    prefix = 'Pad_' if layer == 'Front' else ''
    df = parse_core_shrinkage_content(_content(prefix), layer)
    df['grid_id'] = df['location'].map(GRID)
    return df


# ---------------------------------------------------------------------------
# Filename detection
# ---------------------------------------------------------------------------

def test_parse_cs_filename_front():
    r = parse_cs_filename("C4F_Core_Shrinkage_Panel_28_Front.txt")
    assert r['is_core_shrinkage'] is True
    assert r['Panel Number'] == "28"
    assert r['layer'] == "Front"


def test_parse_cs_filename_back():
    r = parse_cs_filename("C4F_Core_Shrinkage_Panel_28_Back.txt")
    assert r['is_core_shrinkage'] is True
    assert r['layer'] == "Back"


def test_parse_cs_filename_non_cs():
    r = parse_cs_filename("Post PFC_Pad to Via_Panel 20_F.txt")
    assert r['is_core_shrinkage'] is False


# ---------------------------------------------------------------------------
# Content parser
# ---------------------------------------------------------------------------

def test_parse_content_row_count_and_tags():
    df = parse_core_shrinkage_content(_content('Pad_'), 'Front')
    assert len(df) == 16
    assert (df['layer'] == 'Front').all()
    assert (df['Process'] == 'Core Shrinkage').all()


def test_parse_content_strips_pad_prefix():
    df = parse_core_shrinkage_content(_content('Pad_'), 'Front')
    assert 'UL_1' in set(df['location'])
    assert not any(loc.startswith('Pad_') for loc in df['location'])


def test_parse_content_empty_returns_empty_frame():
    assert parse_core_shrinkage_content("no circles here", 'Back').empty


# ---------------------------------------------------------------------------
# Frame builders
# ---------------------------------------------------------------------------

def test_sort_by_grid_anticlockwise_order():
    out = sort_by_grid(_parsed('Front'))
    assert out['grid_id'].tolist() == [
        11, 12, 13, 14, 21, 22, 23, 24, 31, 32, 33, 34, 41, 42, 43, 44]


def test_sort_by_grid_noop_without_grid_column():
    df = parse_core_shrinkage_content(_content(''), 'Back')  # no grid_id added
    assert sort_by_grid(df)['location'].tolist() == df['location'].tolist()


def test_layer_frame_columns():
    assert list(layer_frame(_parsed('Front')).columns) == ['Quarter', 'Grid', 'X', 'Y']


def test_both_frame_columns_and_alignment():
    out = both_frame(_parsed('Front'), _parsed('Back'))
    assert list(out.columns) == ['Quarter', 'Grid', 'Front X', 'Front Y', 'Back X', 'Back Y']
    assert len(out) == 16
    for bad in ('DX', 'DY', 'Distance'):
        assert not any(bad in c for c in out.columns)


def test_both_frame_retains_unmapped_grid():
    front = _parsed('Front')
    back = _parsed('Back')
    front.loc[front['location'] == 'UL_1', 'grid_id'] = np.nan
    back.loc[back['location'] == 'UL_1', 'grid_id'] = np.nan
    out = both_frame(front, back)
    assert len(out) == 16
    assert pd.isna(out['Grid']).sum() == 1


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def _sheets(front_df, back_df, meta=None, warnings=None):
    pytest.importorskip('openpyxl')
    import openpyxl
    from src.excel_export import write_core_shrinkage_excel
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_core_shrinkage_excel(w, front_df, back_df, meta or {}, warnings or [])
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def _sheets_with_qp(front_df, back_df, meta=None, warnings=None):
    pytest.importorskip('openpyxl')
    import openpyxl
    from src.excel_export import write_core_shrinkage_excel
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as w:
        write_core_shrinkage_excel(
            w, front_df, back_df, meta or {}, warnings or [], _NOMINAL)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def test_excel_both_files_sheets():
    wb = _sheets(_parsed('Front'), _parsed('Back'))
    assert wb.sheetnames == [
        'Core Shrinkage - Front', 'Core Shrinkage - Back',
        'Core Shrinkage', 'Report', 'Warnings']


def test_excel_front_only_sheets():
    wb = _sheets(_parsed('Front'), pd.DataFrame())
    assert 'Core Shrinkage - Front' in wb.sheetnames
    assert 'Core Shrinkage - Back' not in wb.sheetnames
    assert 'Core Shrinkage' not in wb.sheetnames
    assert 'Report' in wb.sheetnames


def test_excel_qp_and_sf_sheets_present_both_files():
    wb = _sheets_with_qp(_nominal_df('Front'), _nominal_df('Back'))
    names = wb.sheetnames
    assert 'CS Summary' in names
    assert 'QP Shrinkage - Front' in names
    assert 'QP Shrinkage - Back' in names
    assert 'Scale Factor - Front' in names
    assert 'Scale Factor - Back' in names


def test_excel_qp_front_only():
    wb = _sheets_with_qp(_nominal_df('Front'), pd.DataFrame())
    names = wb.sheetnames
    assert 'QP Shrinkage - Front' in names
    assert 'QP Shrinkage - Back' not in names


def test_excel_without_nominal_coords_no_qp_sheets():
    wb = _sheets(_nominal_df('Front'), _nominal_df('Back'))
    names = wb.sheetnames
    assert not any('QP' in n or 'Scale Factor' in n or 'Summary' in n for n in names)


def test_excel_affine_and_residual_sheets():
    wb = _sheets_with_qp(_nominal_df('Front'), _nominal_df('Back'))
    names = wb.sheetnames
    assert 'Affine - Front' in names and 'Affine - Back' in names
    assert 'Residuals - Front' in names and 'Residuals - Back' in names


def test_excel_residual_sheet_is_um_with_outlier_col():
    """Residual sheet headers are µm and include the Outlier flag column."""
    front = _measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6,
                         outlier='UL_2', outlier_dxy=(0.0, 0.6))
    wb = _sheets_with_qp(front, pd.DataFrame())
    ws = wb['Residuals - Front']
    headers = [c.value for c in ws[1]]
    assert 'Outlier' in headers
    assert any('µm' in str(h) for h in headers)


# ---------------------------------------------------------------------------
# QP Shrinkage
# ---------------------------------------------------------------------------

def test_compute_qp_returns_4_rows():
    qp = compute_qp_shrinkage(_nominal_df('Front'))
    assert len(qp) == 4
    assert list(qp.columns) == ['Quadrant', 'X1', 'X2', 'Y1', 'Y2']
    assert set(qp['Quadrant']) == {'UL', 'LL', 'LR', 'UR'}


def test_compute_qp_zero_when_measured_equals_nominal():
    qp = compute_qp_shrinkage(_nominal_df('Front'))
    for col in ('X1', 'X2', 'Y1', 'Y2'):
        assert (qp[col].abs() < 1e-9).all(), f"Expected 0 in {col}, got {qp[col].tolist()}"


def test_compute_qp_missing_fiducial_returns_nan():
    df = _nominal_df('Front')
    df = df[df['location'] != 'UL_1'].copy()
    qp = compute_qp_shrinkage(df)
    ul_row = qp[qp['Quadrant'] == 'UL'].iloc[0]
    assert pd.isna(ul_row['X1'])
    assert pd.isna(ul_row['Y1'])
    assert pd.notna(ul_row['X2'])
    assert pd.notna(ul_row['Y2'])
    for quad in ('LL', 'LR', 'UR'):
        row = qp[qp['Quadrant'] == quad].iloc[0]
        assert not any(pd.isna(row[c]) for c in ('X1', 'X2', 'Y1', 'Y2'))


def test_compute_qp_empty_input():
    assert compute_qp_shrinkage(pd.DataFrame()).empty


def test_qp_shrinkage_um_units():
    """qp_shrinkage_um is the mm result × 1000 with µm-labelled columns."""
    df = _measured_df('Front', sx=1 - 200e-6, sy=1 - 200e-6)
    mm = compute_qp_shrinkage(df)
    um = qp_shrinkage_um(df)
    assert list(um.columns) == ['Quadrant', 'X1 (µm)', 'X2 (µm)', 'Y1 (µm)', 'Y2 (µm)']
    assert um['X1 (µm)'].iloc[0] == pytest.approx(mm['X1'].iloc[0] * 1000, abs=0.2)


# ---------------------------------------------------------------------------
# Scale Factor
# ---------------------------------------------------------------------------

def test_compute_scale_factor_one_at_nominal():
    sf = compute_scale_factor(_nominal_df('Back'))
    for _, row in sf[sf['Quadrant'] != 'Overall'].iterrows():
        assert abs(row['X Scale'] - 1.0) < 1e-9
        assert abs(row['Y Scale'] - 1.0) < 1e-9
        assert abs(row['X Shrinkage (ppm)']) < 1e-3
        assert abs(row['Y Shrinkage (ppm)']) < 1e-3


def test_compute_scale_factor_overall_row():
    sf = compute_scale_factor(_nominal_df('Front'))
    assert len(sf) == 5
    assert sf.iloc[-1]['Quadrant'] == 'Overall'


def test_compute_scale_factor_empty_input():
    assert compute_scale_factor(pd.DataFrame()).empty


# ---------------------------------------------------------------------------
# Affine Analysis — recovery of a KNOWN transform
# ---------------------------------------------------------------------------

def test_affine_identity_at_nominal():
    analysis = compute_affine_analysis(_nominal_df('Front'), _NOMINAL)
    params = analysis['params'].set_index('Parameter')
    assert abs(params.loc['Rotation', 'Value']) < 1e-6
    assert abs(params.loc['X Scale', 'Value'] - 1.0) < 1e-6
    assert abs(params.loc['Y Scale', 'Value'] - 1.0) < 1e-6
    assert (analysis['residuals']['Res Distance (mm)'].abs() < 1e-9).all()
    assert analysis['pattern'] == 'Stable'


def test_affine_recovers_translation():
    shifted = _measured_df('Front', tx=300.0, ty=325.0)
    params = compute_affine_analysis(shifted, _NOMINAL)['params'].set_index('Parameter')
    assert abs(params.loc['Translation X', 'Value'] - 300.0) < 1e-4
    assert abs(params.loc['Translation Y', 'Value'] - 325.0) < 1e-4
    assert abs(params.loc['Rotation', 'Value']) < 1e-4


def test_affine_recovers_rotation():
    rotated = _measured_df('Back', theta_deg=0.1, tx=300.0, ty=320.0)
    analysis = compute_affine_analysis(rotated, _NOMINAL)
    params = analysis['params'].set_index('Parameter')
    assert abs(params.loc['Rotation', 'Value'] - 0.1) < 1e-4
    assert (analysis['residuals']['Res Distance (mm)'].abs() < 1e-6).all()


def test_affine_recovers_anisotropic_scale_with_axis_identity():
    """QR decomposition keeps X-scale on X (SVD could swap). Sx≠Sy recovered."""
    sx, sy = 1 - 700e-6, 1 - 200e-6
    df = _measured_df('Front', sx=sx, sy=sy, tx=300.0, ty=320.0)
    params = compute_affine_analysis(df, _NOMINAL)['params'].set_index('Parameter')
    assert params.loc['X Shrinkage', 'Value'] == pytest.approx(-700.0, abs=2.0)
    assert params.loc['Y Shrinkage', 'Value'] == pytest.approx(-200.0, abs=2.0)


def test_affine_pattern_uniform_shrinkage():
    df = _measured_df('Front', sx=1 - 200e-6, sy=1 - 200e-6)
    assert compute_affine_analysis(df, _NOMINAL)['pattern'] == 'Uniform shrinkage'


def test_affine_pattern_anisotropic():
    df = _measured_df('Front', sx=1 - 700e-6, sy=1 - 200e-6)
    assert compute_affine_analysis(df, _NOMINAL)['pattern'] == 'Anisotropic shrinkage'


def test_affine_pattern_outlier():
    df = _measured_df('Front', sx=1 - 200e-6, sy=1 - 200e-6,
                      outlier='UL_2', outlier_dxy=(0.0, 0.6))
    assert 'outlier' in compute_affine_analysis(df, _NOMINAL)['pattern'].lower()


def test_affine_residuals_shape():
    analysis = compute_affine_analysis(_nominal_df('Front'), _NOMINAL)
    res = analysis['residuals']
    assert len(res) == 16
    assert list(res.columns) == [
        'location', 'Quadrant', 'Res X (mm)', 'Res Y (mm)', 'Res Distance (mm)']


# ---------------------------------------------------------------------------
# Shrinkage summary + outlier robustness
# ---------------------------------------------------------------------------

def test_summary_recovers_known_scale_no_outlier():
    df = _measured_df('Back', sx=1 - 465e-6, sy=1 - 465e-6, tx=300.0, ty=320.0)
    s = compute_shrinkage_summary(df, _NOMINAL)
    assert s['n_outliers'] == 0
    assert s['Sx_ppm'] == pytest.approx(-465.0, abs=2.0)
    assert s['Sy_ppm'] == pytest.approx(-465.0, abs=2.0)
    # With no outliers the robust refit equals the all-dot fit.
    assert s['Sx_ppm_robust'] == pytest.approx(s['Sx_ppm'], abs=1e-6)


def test_summary_flags_outlier_and_robust_recovers_truth():
    df = _measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6, tx=300.0, ty=320.0,
                      outlier='UL_2', outlier_dxy=(0.0, 0.6))
    s = compute_shrinkage_summary(df, _NOMINAL)
    assert s['n_outliers'] == 1 and s['outliers'] == ['UL_2']
    # Robust (outlier-excluded) recovers the true scale exactly; the all-dot fit
    # is pulled away by the bad dot, so the two disagree.
    assert s['Sy_ppm_robust'] == pytest.approx(-465.0, abs=1.0)
    assert abs(s['Sy_ppm'] - s['Sy_ppm_robust']) > 2.0
    assert s['max_residual_um'] > 50.0


def test_degenerate_stuck_probe_refuses_headline():
    """A collapsed axis (stuck probe) must yield no analysis, not ±1e6 ppm."""
    df = _measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6)
    df['coord_x'] = 7.0
    assert analyze_layer(df, _NOMINAL) == {}
    assert compute_affine_analysis(df, _NOMINAL) == {}
    assert compute_shrinkage_summary(df, _NOMINAL) == {}


def test_degenerate_scale_boundary():
    """A halved panel still reports; a 99.5%-collapsed axis is refused."""
    ok_df = _measured_df('Front', sx=0.5, sy=0.5)
    assert analyze_layer(ok_df, _NOMINAL)['n_points'] == 16
    bad_df = _measured_df('Front', sx=0.005, sy=1.0)
    assert analyze_layer(bad_df, _NOMINAL) == {}


def test_numbering_swap_keeps_its_diagnostic():
    """A numbering swap (Sy≈0.03, huge residuals) is mismatched data, not a
    collapse: analysis must stay available so the swap is diagnosable."""
    good = _measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6, tx=300.0, ty=320.0)
    bad = _swap_pairs(good)
    assert analyze_layer(bad, _NOMINAL)['n_points'] == 16


def test_flag_outliers_marks_exactly_the_bad_dot():
    df = _measured_df('Front', outlier='LR_3', outlier_dxy=(0.5, 0.0))
    res = flag_outliers(compute_affine_analysis(df, _NOMINAL), outlier_um=50.0)
    flagged = set(res[res['Outlier']]['location'])
    assert flagged == {'LR_3'}


def test_summary_frame_columns():
    s = compute_shrinkage_summary(_measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6),
                                  _NOMINAL)
    sf = summary_frame(s)
    assert list(sf.columns) == ['Metric', 'Value']
    assert 'X Shrinkage (ppm)' in set(sf['Metric'])


def test_affine_residuals_um_has_outlier_and_um_cols():
    df = _measured_df('Front', outlier='UL_2', outlier_dxy=(0.0, 0.6))
    res = affine_residuals_um(compute_affine_analysis(df, _NOMINAL), outlier_um=50.0)
    assert 'Res Dist (µm)' in res.columns
    assert 'Outlier' in res.columns
    assert res['Outlier'].sum() == 1


# ---------------------------------------------------------------------------
# Front vs Back registration
# ---------------------------------------------------------------------------

def test_front_back_registration_difference_um():
    front = _measured_df('Front', tx=300.0, ty=320.0)
    back  = _measured_df('Back',  tx=300.5, ty=320.0)  # 0.5 mm shift in X
    fb = compute_front_back_registration(front, back)
    assert list(fb.columns) == ['Quarter', 'Grid', 'ΔX (µm)', 'ΔY (µm)', 'Δ Dist (µm)']
    assert len(fb) == 16
    # Front − Back ⇒ −500 µm in X everywhere.
    assert fb['ΔX (µm)'].round().eq(-500.0).all()
    assert fb['ΔY (µm)'].abs().max() < 1e-3


def test_front_back_registration_needs_both_layers():
    assert compute_front_back_registration(_measured_df('Front'), pd.DataFrame()).empty


# ---------------------------------------------------------------------------
# Centered deviation table + analyze_layer single-pass
# ---------------------------------------------------------------------------

def test_centered_table_matches_deviation_field():
    df = _measured_df('Back', sx=1 - 465e-6, sy=1 - 465e-6, theta_deg=0.05,
                      tx=300.0, ty=320.0)
    tbl = compute_centered_deviation_table(df, _NOMINAL)
    dev = compute_deviation_field(df, _NOMINAL, field='total')
    rows = tbl[tbl['location'] != 'SUMMARY (RMS)']
    assert len(rows) == 16
    # No broken per-coordinate ppm columns
    assert not any('PPM' in c.upper() for c in tbl.columns)
    # Δ Dist matches the Distortion Map's total-deviation magnitude per dot
    for loc in ('UL_1', 'LR_3', 'UR_2'):
        t = rows[rows['location'] == loc]['Δ Dist (µm)'].iloc[0]
        d = dev[dev['location'] == loc]['mag_um'].iloc[0]
        assert t == pytest.approx(d, abs=0.2)


def test_centered_table_has_summary_rms_row():
    df = _measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6, tx=300.0, ty=320.0)
    tbl = compute_centered_deviation_table(df, _NOMINAL)
    last = tbl.iloc[-1]
    assert last['location'] == 'SUMMARY (RMS)'
    rows = tbl[tbl['location'] != 'SUMMARY (RMS)']
    expect = np.sqrt(np.mean(rows['Δ Dist (µm)'].astype(float) ** 2))
    assert float(last['Δ Dist (µm)']) == pytest.approx(expect, abs=0.2)


def test_centered_table_coords_reconstruct_delta():
    """Centered − CAM nominal (mm) equals the reported Δ (µm)."""
    df = _measured_df('Back', sx=1 - 300e-6, sy=1 - 300e-6, tx=300.0, ty=320.0)
    rows = compute_centered_deviation_table(df, _NOMINAL)
    rows = rows[rows['location'] != 'SUMMARY (RMS)']
    dx = (rows['Centered X'].astype(float) - rows['CAM Nominal X'].astype(float)) * 1000.0
    assert np.allclose(dx.values, rows['ΔX (µm)'].astype(float).values, atol=0.2)


def test_analyze_layer_single_pass_matches_individual_calls():
    df = _measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6, tx=300.0, ty=320.0,
                      outlier='UL_2', outlier_dxy=(0.0, 0.6))
    a = analyze_layer(df, _NOMINAL)
    assert a and a['n_points'] == 16 and a['outliers'] == ['UL_2']
    # Passing the precomputed analysis gives identical results to computing fresh.
    s1 = compute_shrinkage_summary(df, _NOMINAL)
    s2 = compute_shrinkage_summary(df, _NOMINAL, analysis=a)
    assert s1 == s2
    p1 = compute_affine_analysis(df, _NOMINAL)['pattern']
    p2 = compute_affine_analysis(df, _NOMINAL, analysis=a)['pattern']
    assert p1 == p2


def test_analyze_layer_empty_below_four():
    assert analyze_layer(_measured_df('Front').head(3), _NOMINAL) == {}


# ---------------------------------------------------------------------------
# Regression: the −471 mm numbering-swap failure mode must stay loud
# ---------------------------------------------------------------------------

def _swap_pairs(df: pd.DataFrame) -> pd.DataFrame:
    """Swap within-quadrant point numbering 1↔2, 3↔4 (the QVM-vs-CAM mismatch)."""
    swap = {1: 2, 2: 1, 3: 4, 4: 3}
    out = df.copy()
    out['point_num'] = out['point_num'].map(swap)
    out['location'] = out['quadrant'] + '_' + out['point_num'].astype(str)
    return out


# ---------------------------------------------------------------------------
# Deviation field (Distortion Map) + rigid fit + direction classification
# ---------------------------------------------------------------------------

def test_rigid_fit_recovers_rotation_and_keeps_scale():
    # Expanded + rotated panel: rigid fit recovers rotation, scale stays in the residual.
    df = _measured_df('Front', sx=1.001, sy=1.001, theta_deg=0.1, tx=300.0, ty=320.0)
    nom, meas, _, _ = (None, None, None, None)
    from src.core_shrinkage_parser import _matched_arrays
    nom, meas, _, _ = _matched_arrays(df, _NOMINAL)
    R, t = _rigid_fit(nom, meas)
    theta = np.degrees(np.arctan2(R[1, 0], R[0, 0]))
    assert theta == pytest.approx(0.1, abs=1e-3)
    resid = meas - ((R @ nom.T).T + t)
    assert np.linalg.norm(resid, axis=1).max() > 0.05  # ~mm-scale expansion remains


def test_deviation_field_total_shows_expansion():
    df = _measured_df('Front', sx=1.0005, sy=1.0005, tx=300.0, ty=320.0)
    dev = compute_deviation_field(df, _NOMINAL, field='total')
    assert len(dev) == 16
    non_stable = dev[dev['klass'] != 'stable']
    assert (non_stable['klass'] == 'expansion').all()
    assert (dev['klass'] == 'expansion').sum() >= 12


def test_deviation_field_total_shows_contraction_when_shrunk():
    df = _measured_df('Front', sx=1 - 500e-6, sy=1 - 500e-6, tx=300.0, ty=320.0)
    dev = compute_deviation_field(df, _NOMINAL, field='total')
    non_stable = dev[dev['klass'] != 'stable']
    assert (non_stable['klass'] == 'contraction').all()


def test_deviation_field_residual_clean_is_small():
    df = _measured_df('Back', sx=1 - 465e-6, sy=1 - 465e-6, tx=300.0, ty=320.0)
    dev = compute_deviation_field(df, _NOMINAL, field='residual')
    assert dev['mag_um'].max() < 1.0          # scale removed ⇒ ~0


def test_deviation_field_marks_outlier():
    df = _measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6,
                      outlier='UL_2', outlier_dxy=(0.0, 0.6))
    dev = compute_deviation_field(df, _NOMINAL, field='residual')
    row = dev[dev['location'] == 'UL_2'].iloc[0]
    assert row['outlier'] and row['klass'] == 'outlier'
    assert (dev[dev['location'] != 'UL_2']['klass'] != 'outlier').all()


def test_classify_vector_directions():
    c = (0.0, 0.0)
    assert _classify_vector(10, 0, 100, 0, c) == 'expansion'    # outward
    assert _classify_vector(-10, 0, 100, 0, c) == 'contraction'  # inward
    assert _classify_vector(0, 10, 100, 0, c) == 'twist'          # tangential
    assert _classify_vector(0.1, 0, 100, 0, c) == 'stable'        # below threshold


def test_deviation_field_empty_below_four():
    df = _measured_df('Front').head(3)
    assert compute_deviation_field(df, _NOMINAL).empty


def test_swap_regression_produces_large_residual():
    good = _measured_df('Front', sx=1 - 465e-6, sy=1 - 465e-6, tx=300.0, ty=320.0)
    bad = _swap_pairs(good)

    good_a = compute_affine_analysis(good, _NOMINAL)
    bad_a = compute_affine_analysis(bad, _NOMINAL)

    good_max = good_a['residuals']['Res Distance (mm)'].abs().max()
    bad_max = bad_a['residuals']['Res Distance (mm)'].abs().max()

    assert good_max < 0.001          # correctly ordered ⇒ microns
    assert bad_max > 100.0           # swapped ⇒ catastrophic (the −471 mm bug)
    assert 'outlier' in bad_a['pattern'].lower() or 'distortion' in bad_a['pattern'].lower()


# ---------------------------------------------------------------------------
# Shipped configuration
# ---------------------------------------------------------------------------

def test_shipped_nominal_coords_match_the_test_fixture(settings):
    """_NOMINAL above is a copy of CORE_SHRINKAGE.NOMINAL_COORDS. If someone
    retunes the shipped panel geometry, every expectation in this file silently
    stops describing the real config — so pin them together."""
    shipped = settings['CORE_SHRINKAGE']['NOMINAL_COORDS']
    assert set(shipped) == set(_NOMINAL)
    for site, nominal in _NOMINAL.items():
        assert shipped[site]['x'] == pytest.approx(nominal['x']), site
        assert shipped[site]['y'] == pytest.approx(nominal['y']), site


def test_shipped_nominal_coords_cover_every_grid_site(settings):
    """A fiducial present in GRID_MAPPING but missing a nominal coordinate is
    dropped from every affine fit, silently shrinking the sample."""
    assert set(settings['CORE_SHRINKAGE']['NOMINAL_COORDS']) == set(settings['GRID_MAPPING'])


def test_shipped_residual_zones_are_ordered(settings):
    """The Distortion Map colours by these bands; out-of-order values would make
    a band unreachable."""
    zones = settings['CORE_SHRINKAGE']['residual_zones_um']
    values = [zones['exceptional'], zones['excellent'], zones['good'], zones['acceptable']]
    assert values == sorted(values)
    assert values[-1] <= settings['CORE_SHRINKAGE']['outlier_um']


# ---------------------------------------------------------------------------
# Sign convention lock (Finding 1)
# ---------------------------------------------------------------------------

def test_shrinkage_sign_convention_agrees_across_modules():
    """Shrink = NEGATIVE ppm on both paths (measured − nominal).

    Convention lock against Finding 1: a uniformly shrunk panel must read
    negative on the Core Shrinkage scale-factor path AND on the registration
    distance path, with matching magnitudes — the customer reference
    workbook (Core Shrinkage Calculation.xlsx) uses measured − design.
    """
    from src.registration_shrinkage import compute_distance_shrinkage
    df = _measured_df('Front', sx=1 - 500e-6, sy=1 - 500e-6)
    sf = compute_scale_factor(df)
    per_quad = sf[sf['Quadrant'] != 'Overall']
    assert (per_quad['X Shrinkage (ppm)'] < 0).all()
    assert (per_quad['Y Shrinkage (ppm)'] < 0).all()
    reg = compute_distance_shrinkage(
        pd.DataFrame({'Location': df['location'],
                      'Coord. X': df['coord_x'],
                      'Coord. Y': df['coord_y']}), _NOMINAL)
    assert not reg.quadrant.empty
    assert (reg.quadrant['QP Shrinkage X (PPM)'] < 0).all()
    assert (reg.quadrant['QP Shrinkage Y (PPM)'] < 0).all()
    assert abs(per_quad['X Shrinkage (ppm)'].mean()
               - reg.quadrant['QP Shrinkage X (PPM)'].mean()) < 2.0
