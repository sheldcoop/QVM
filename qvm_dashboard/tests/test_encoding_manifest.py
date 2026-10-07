"""Encoding fallback + workbook file manifest (Finding 2).

Uploads may arrive in UTF-8, UTF-8-with-BOM, Latin-1, or UTF-16 (common on
Windows metrology stations). Every exported workbook must carry a Files
manifest so a reviewer can prove which uploads are inside — including
files that were rejected.
"""

from io import BytesIO
from types import SimpleNamespace

import pandas as pd
import pytest

from src.utils import decode_upload_bytes

VALID_CONTENT = """
Circle: Pad_LL_2(ID:1, From 180 Pts.)
     Coord. X =        63.857203
     Coord. Y =       104.634501
     Diameter =         0.308911

Circle: Via_LL_2(ID:2, From 70 Pts.)
     Coord. X =        63.856611
     Coord. Y =       104.630820
     Diameter =         0.016742

Distance: PtV_LL_2(ID:3) between Via_LL_2(ID:2) and Pad_LL_2(ID:1)
           SC =         0.003728
           DX =         0.000592
           DY =         0.003681
"""


def _upload(name, data):
    return SimpleNamespace(name=name, getvalue=lambda: data)


# ---------------------------------------------------------------------------
# Decoding
# ---------------------------------------------------------------------------


def test_decode_plain_utf8():
    text, enc = decode_upload_bytes(VALID_CONTENT.encode("utf-8"))
    assert enc == "utf-8"
    assert "Circle" in text


def test_decode_utf8_bom_strips_marker():
    text, enc = decode_upload_bytes(b"\xef\xbb\xbf" + VALID_CONTENT.encode("utf-8"))
    assert enc == "utf-8-sig"
    assert not text.startswith("\ufeff")


def test_decode_latin1():
    data = "Remarque: pi\u00e8ce\n".encode("latin-1") + VALID_CONTENT.encode("latin-1")
    text, enc = decode_upload_bytes(data)
    assert enc == "latin-1"
    assert "pi\u00e8ce" in text


def test_decode_utf16():
    text, enc = decode_upload_bytes(VALID_CONTENT.encode("utf-16"))
    assert enc == "utf-16"
    assert "Circle" in text


def test_decode_empty_is_utf8():
    text, enc = decode_upload_bytes(b"")
    assert (text, enc) == ("", "utf-8")


# ---------------------------------------------------------------------------
# Ingestion records the detected encoding
# ---------------------------------------------------------------------------


def test_ingest_utf16_content_parses_and_records_encoding(settings):
    from src.ingestion import ingest_uploaded_file

    df, _, _, _, finfo = ingest_uploaded_file(
        _upload("PTV_Panel_20_F.txt", VALID_CONTENT.encode("utf-16")), settings
    )
    assert finfo["detected_encoding"] == "utf-16"
    assert len(df) == 1 and df["Process"].iloc[0] == "Pad to Via"


def test_ingest_latin1_content_parses_and_records_encoding(settings):
    from src.ingestion import ingest_uploaded_file

    data = "Remarque: pi\u00e8ce\n".encode("latin-1") + VALID_CONTENT.encode("latin-1")
    df, _, _, _, finfo = ingest_uploaded_file(
        _upload("PTV_Panel_20_F.txt", data), settings
    )
    assert finfo["detected_encoding"] == "latin-1"
    assert len(df) == 1


# ---------------------------------------------------------------------------
# Files manifest in both workbooks
# ---------------------------------------------------------------------------


def _records():
    return [
        {
            "filename": "PTV_Panel_20_F.txt",
            "pipeline": "QVM",
            "encoding": "utf-16",
            "rows": 1,
            "status": "ok",
            "detail": "",
        },
        {
            "filename": "broken.txt",
            "pipeline": "QVM",
            "encoding": "utf-8",
            "rows": 0,
            "status": "skipped",
            "detail": "QVMParseError: nope",
        },
        {
            "filename": "=1+1.txt",
            "pipeline": "QVM",
            "encoding": "utf-8",
            "rows": 0,
            "status": "skipped",
            "detail": "rejected",
        },
    ]


def _files_sheet(wb):
    assert "Files" in wb.sheetnames
    ws = wb["Files"]
    header = [c.value for c in ws[1]]
    assert header == ["Filename", "Pipeline", "Encoding", "Rows", "Status", "Detail"]
    return ws


def test_unified_excel_manifest_lists_ok_and_skipped(settings):
    pytest.importorskip("openpyxl")
    import openpyxl
    from src.excel_export import write_unified_excel

    cn = settings["COLUMN_NAMES"]
    df = pd.DataFrame(
        {
            cn["grid_id"]: [11],
            cn["location"]: ["p0"],
            cn["inner_diameter"]: [0.02],
            cn["outer_diameter"]: [0.3125],
            cn["coord_x"]: [65.0],
            cn["coord_y"]: [104.0],
            cn["x_distance"]: [0.001],
            cn["y_distance"]: [0.002],
            cn["ptv_distance"]: [0.003],
            "Build Up": ["BU-01"],
            "Side": ["F"],
            "Panel": ["20"],
            "Process": ["Pad to Via"],
        }
    )
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        write_unified_excel(
            w, df, settings, {"panel": "20", "file_records": _records()}
        )
    buf.seek(0)
    ws = _files_sheet(openpyxl.load_workbook(buf))
    assert ws.max_row == 4
    assert ws.cell(row=3, column=1).value == "broken.txt"
    assert ws.cell(row=3, column=5).value == "skipped"


def test_unified_excel_without_records_has_no_files_sheet(settings):
    pytest.importorskip("openpyxl")
    import openpyxl
    from src.excel_export import write_unified_excel

    cn = settings["COLUMN_NAMES"]
    df = pd.DataFrame(
        {
            cn["grid_id"]: [11],
            cn["location"]: ["p0"],
            cn["inner_diameter"]: [0.02],
            cn["outer_diameter"]: [0.3125],
            cn["coord_x"]: [65.0],
            cn["coord_y"]: [104.0],
            cn["x_distance"]: [0.001],
            cn["y_distance"]: [0.002],
            cn["ptv_distance"]: [0.003],
            "Build Up": ["BU-01"],
            "Side": ["F"],
            "Panel": ["20"],
            "Process": ["Pad to Via"],
        }
    )
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        write_unified_excel(w, df, settings, {"panel": "20"})
    buf.seek(0)
    assert "Files" not in openpyxl.load_workbook(buf).sheetnames


def test_cs_excel_manifest_lists_ok_and_skipped(settings):
    pytest.importorskip("openpyxl")
    import openpyxl
    from src.excel_export import write_core_shrinkage_excel

    front = pd.DataFrame(
        [
            {
                "location": "UL_1",
                "quadrant": "UL",
                "point_num": 1,
                "coord_x": 1.0,
                "coord_y": 2.0,
                "diameter": 0.37,
                "pts_count": 150,
                "layer": "Front",
                "Panel": "28",
            }
        ]
    )
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        write_core_shrinkage_excel(
            w, front, pd.DataFrame(), {"panel": "28", "file_records": _records()}
        )
    buf.seek(0)
    ws = _files_sheet(openpyxl.load_workbook(buf))
    assert ws.max_row == 4


def test_manifest_neutralises_formula_injection(settings):
    """A filename starting with '=' must round-trip as text, not a formula."""
    pytest.importorskip("openpyxl")
    import openpyxl
    from src.excel_export import write_unified_excel

    cn = settings["COLUMN_NAMES"]
    df = pd.DataFrame(
        {
            cn["grid_id"]: [11],
            cn["location"]: ["p0"],
            cn["inner_diameter"]: [0.02],
            cn["outer_diameter"]: [0.3125],
            cn["coord_x"]: [65.0],
            cn["coord_y"]: [104.0],
            cn["x_distance"]: [0.001],
            cn["y_distance"]: [0.002],
            cn["ptv_distance"]: [0.003],
            "Build Up": ["BU-01"],
            "Side": ["F"],
            "Panel": ["20"],
            "Process": ["Pad to Via"],
        }
    )
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        write_unified_excel(
            w, df, settings, {"panel": "20", "file_records": _records()}
        )
    buf.seek(0)
    ws = openpyxl.load_workbook(buf)["Files"]
    evil = ws.cell(row=4, column=1).value
    assert evil != "=1+1.txt"
    assert "1+1" in evil
