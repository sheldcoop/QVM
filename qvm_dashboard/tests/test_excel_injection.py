"""Excel formula-injection neutralisation (Finding 3).

Sidebar text (project/part/lot, Build Up overrides) and filenames flow into
workbook cells. Values starting with = + - @ must round-trip as displayed
text, never as live formulas, in both writers.
"""

from io import BytesIO

import pandas as pd
import pytest

EVIL_META = {
    "panel": "20",
    "project_name": "=1+1",
    "part_number": "+2+2",
    "lot_number": "@SUM(1,1)",
    "lot_name": '=HYPERLINK("http://x")',
}


def _qvm_frame(settings, location="p0"):
    cn = settings["COLUMN_NAMES"]
    return pd.DataFrame(
        {
            cn["grid_id"]: [11],
            cn["location"]: [location],
            cn["inner_diameter"]: [0.02],
            cn["outer_diameter"]: [0.3125],
            cn["coord_x"]: [65.0],
            cn["coord_y"]: [104.0],
            cn["x_distance"]: [0.001],
            cn["y_distance"]: [0.002],
            cn["ptv_distance"]: [0.003],
            "Build Up": ["BU-01"],
            "Side": ["F"],
            "Panel": ["20"],
            "Process": ["Pad to Via"],
        }
    )


def _unified_wb(settings, df, meta):
    pytest.importorskip("openpyxl")
    import openpyxl
    from src.excel_export import write_unified_excel

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        write_unified_excel(w, df, settings, meta)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def _text_cells(ws):
    """All (coordinate, value, data_type) for string-ish cells below header."""
    out = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        for c in row:
            if isinstance(c.value, str):
                out.append((c.coordinate, c.value, c.data_type))
    return out


def test_unified_report_meta_is_text_not_formula(settings):
    ws = _unified_wb(settings, _qvm_frame(settings), dict(EVIL_META))["Report"]
    cells = _text_cells(ws)
    assert cells, "expected string cells in Report"
    for coord, value, dtype in cells:
        assert dtype != "f", f"{coord} stored as formula: {value!r}"
    values = [v for _, v, _ in cells]
    assert any("1+1" in v for v in values)


def test_unified_location_and_buildup_are_text(settings):
    df = _qvm_frame(settings, location="=2+2")
    df["Build Up"] = ["@evil"]
    wb = _unified_wb(settings, df, {"panel": "20"})
    for sheet in ("Pad to Via", "Report"):
        for coord, value, dtype in _text_cells(wb[sheet]):
            assert dtype != "f", f"{sheet}{coord} stored as formula: {value!r}"


def test_unified_warnings_are_text(settings):
    from src.excel_export import write_unified_excel

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        write_unified_excel(
            w,
            _qvm_frame(settings),
            settings,
            {"panel": "20"},
            warnings=["=cmd|evil", "see =1+1 file"],
        )
    buf.seek(0)
    import openpyxl

    ws = openpyxl.load_workbook(buf)["Warnings"]
    for coord, value, dtype in _text_cells(ws):
        assert dtype != "f", f"{coord} stored as formula: {value!r}"


def test_cs_report_meta_is_text_not_formula(settings):
    pytest.importorskip("openpyxl")
    import openpyxl
    from src.excel_export import write_core_shrinkage_excel

    front = pd.DataFrame(
        [
            {
                "location": "UL_1",
                "quadrant": "UL",
                "point_num": 1,
                "coord_x": 1.0,
                "coord_y": 2.0,
                "diameter": 0.37,
                "pts_count": 150,
                "layer": "Front",
                "Panel": "28",
            }
        ]
    )
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        write_core_shrinkage_excel(
            w, front, pd.DataFrame(), {"panel": "28", **EVIL_META}
        )
    buf.seek(0)
    ws = openpyxl.load_workbook(buf)["Report"]
    cells = _text_cells(ws)
    assert cells, "expected string cells in CS Report"
    for coord, value, dtype in cells:
        assert dtype != "f", f"{coord} stored as formula: {value!r}"


def test_cs_warnings_are_text(settings):
    pytest.importorskip("openpyxl")
    import openpyxl
    from src.excel_export import write_core_shrinkage_excel

    front = pd.DataFrame(
        [
            {
                "location": "UL_1",
                "quadrant": "UL",
                "point_num": 1,
                "coord_x": 1.0,
                "coord_y": 2.0,
                "diameter": 0.37,
                "pts_count": 150,
                "layer": "Front",
                "Panel": "28",
            }
        ]
    )
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        write_core_shrinkage_excel(
            w, front, pd.DataFrame(), {"panel": "28"}, ["=evil warning"]
        )
    buf.seek(0)
    ws = openpyxl.load_workbook(buf)["Warnings"]
    for coord, value, dtype in _text_cells(ws):
        assert dtype != "f", f"{coord} stored as formula: {value!r}"


def test_normal_text_untouched(settings):
    """Ordinary values must not gain stray apostrophes."""
    ws = _unified_wb(
        settings, _qvm_frame(settings), {"panel": "20", "project_name": "C4F line 3"}
    )["Report"]
    values = [
        c.value
        for row in ws.iter_rows(min_row=2, values_only=False)
        for c in row
        if isinstance(c.value, str)
    ]
    assert "C4F line 3" in values
    assert not any(v.startswith("'") for v in values)


def test_text_columns_are_sanitised_whatever_their_dtype(settings):
    """A Build Up or Side starting with "=" must never become a formula.

    pandas 3 gives text columns the ``str`` dtype instead of ``object``; the
    sanitiser used to check for ``object`` only and silently skipped them.
    """
    from src.excel_export._shared import _sanitize_frame

    df = _qvm_frame(settings)
    df["Build Up"] = ["=1+1"]
    df["Side"] = ['=HYPERLINK("http://x","F")']
    for column in ("Build Up", "Side"):
        for dtype in (df[column].dtype, object):
            cleaned = _sanitize_frame(df.astype({column: dtype}))
            assert not str(cleaned[column].iloc[0]).startswith("="), (column, dtype)

    wb = _unified_wb(settings, df, {"panel": "20"})
    for sheet in ("Pad to Via", "Report", "Shrinkage Summary"):
        formulas = [
            (c.coordinate, c.value)
            for row in wb[sheet].iter_rows()
            for c in row
            if c.data_type == "f"
        ]
        assert formulas == [], f"{sheet}: {formulas}"
