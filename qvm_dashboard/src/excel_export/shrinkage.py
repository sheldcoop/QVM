"""Registration shrinkage workbook (the Shrinkage tab's own export).

Sheets:
- ``Summary``           one line per Build Up / Panel / Side
- ``Quadrants``         quadrant values of every group
- ``Quadrant Sizes``    measured against design size (also for groups without a result)
- ``Calculation``       every step as live Excel formulas, from the positions to the panel value
- ``Charts``            the panel map and the quadrant bars per group
- ``Notes``             method, measured point, design notes and warnings
"""

import logging

import pandas as pd

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.registration_shrinkage import (
    METHOD_LABEL,
    _point_map,
    compute_shrinkage_by_group,
    shrinkage_config,
    summary_frame,
)
from src.visuals.shrinkage import plot_quadrant_bars, plot_shrinkage_panel

from ._shared import WarnFn, _fig_to_xl_image, _notify, _safe_text, _sanitize_frame, _write_chart_fallback

logger = logging.getLogger(__name__)

_CHART_WIDTH, _CHART_HEIGHT = 520, 380
_ROWS_PER_CHART = 21


def _stack(groups, pick) -> pd.DataFrame:
    """Concatenate one frame per group, each prefixed with the group's labels."""
    frames = []
    for group in groups:
        frame = pick(group)
        if frame is None or frame.empty:
            continue
        frame = frame.copy()
        for position, (name, value) in enumerate(group.group.items()):
            frame.insert(position, name, value)
        frames.append(frame)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def _write(writer: pd.ExcelWriter, name: str, frame: pd.DataFrame, empty_note: str) -> None:
    if frame.empty:
        frame = pd.DataFrame({'info': [empty_note]})
    frame.to_excel(writer, sheet_name=name, index=False)
    _copper_header(writer.sheets[name])


_QUADRANTS = ('UL', 'LL', 'LR', 'UR')
_BOLD = Font(bold=True)
_HEAD = PatternFill('solid', start_color='B87333')            # copper, as in the app
_HEAD_FONT = Font(bold=True, color='FFFFFF')
_RESULT = PatternFill('solid', start_color='F6E3D0')
_EXPLAIN = (
    'How to read this sheet. Every number below the positions is an Excel formula: click a cell to see it. '
    'Change a Design value (shaded) and everything recalculates.',
    'Points of a quadrant: 1 = top-left, 2 = bottom-left, 3 = bottom-right, 4 = top-right.',
    'X1 = x(4) - x(1), the top edge.  X2 = x(3) - x(2), the bottom edge.  X length = average of X1 and X2.',
    'Y1 = y(1) - y(2), the left edge.  Y2 = y(4) - y(3), the right edge.  Y length = average of Y1 and Y2.',
    'Delta (um) = (measured length - design length) x 1000.  This is how much one quadrant changed, in micrometres.',
    'ppm = (measured length / design length - 1) x 1,000,000.  The same change, relative to the length.',
    'Panel = average of the four quadrants.  X - Y = panel X ppm minus panel Y ppm.  '
    'Quadrant range = largest minus smallest quadrant ppm.',
    'Negative = shrink, positive = grown.',
)
_STEP_HEADERS = (
    'Quadrant',
    'X design (mm)', 'X1 measured', 'X2 measured', 'X measured (mm)',
    'Delta X (um)', 'Shrinkage X (ppm)',
    'Y design (mm)', 'Y1 measured', 'Y2 measured', 'Y measured (mm)',
    'Delta Y (um)', 'Shrinkage Y (ppm)',
)


def _header(sheet, row: int, titles) -> None:
    for column, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.fill, cell.font = _HEAD, _HEAD_FONT


def _copper_header(sheet) -> None:
    """Copper header row and readable column widths for a table sheet."""
    for cell in sheet[1]:
        cell.fill, cell.font = _HEAD, _HEAD_FONT
        width = max(len(str(cell.value or '')) + 3, 12)
        sheet.column_dimensions[cell.column_letter].width = min(width, 60)
    sheet.freeze_panes = 'A2'


def _write_calculation(sheet, groups) -> None:
    """The company calculation as live formulas, one block per reported group.

    Only the design quadrant size and the measured positions are typed in as
    numbers; lengths, um, ppm, panel average, X - Y and range are formulas on
    them, so the user can follow and check every step, or change the design
    size and watch the result move. The app's own result is printed below for
    comparison. Formulas are built from cell addresses only, never from text
    of the file.
    """
    for row, line in enumerate(_EXPLAIN, start=1):
        sheet.cell(row=row, column=1, value=line)
    sheet.cell(row=1, column=1).font = _BOLD
    row = len(_EXPLAIN) + 2

    written = False
    for group in groups:
        if group.result.quadrant.empty or not group.lengths:
            continue
        measured = _point_map(group.measured, *group.columns)
        quadrants = [q for q in _QUADRANTS if all(f'{q}_{c}' in measured for c in '1234')]
        if not quadrants:
            continue
        written = True

        sheet.cell(row=row, column=1,
                   value=_safe_text(f'{group.process} | {group.label} | {group.point} centres')).font = _BOLD
        row += 1
        sheet.cell(row=row, column=1, value='Step 1: design quadrant size and measured positions (mm)').font = _BOLD
        row += 1
        design_x, design_y = f'$B${row}', f'$B${row + 1}'
        for title, value in (('Design X (mm)', group.lengths[0]), ('Design Y (mm)', group.lengths[1])):
            sheet.cell(row=row, column=1, value=title).font = _BOLD
            sheet.cell(row=row, column=2, value=float(value)).fill = _RESULT
            row += 1
        row += 1
        _header(sheet, row, ('Point', 'Measured X', 'Measured Y'))
        at = {}  # point name -> its row
        for quadrant in quadrants:
            for corner in '1234':
                row += 1
                name = f'{quadrant}_{corner}'
                at[name] = row
                for column, value in enumerate(
                        (name, measured[name]['x'], measured[name]['y']), start=1):
                    sheet.cell(row=row, column=column, value=value)

        row += 2
        sheet.cell(row=row, column=1, value='Step 2: lengths, change in um, change in ppm').font = _BOLD
        row += 1
        _header(sheet, row, _STEP_HEADERS)
        first = row + 1
        for quadrant in quadrants:
            row += 1
            p1, p2, p3, p4 = (at[f'{quadrant}_{c}'] for c in '1234')
            cells = (
                quadrant,
                f'={design_x}', f'=B{p4}-B{p1}', f'=B{p3}-B{p2}', f'=AVERAGE(C{row}:D{row})',
                f'=(E{row}-B{row})*1000', f'=(E{row}/B{row}-1)*1000000',
                f'={design_y}', f'=C{p1}-C{p2}', f'=C{p4}-C{p3}', f'=AVERAGE(I{row}:J{row})',
                f'=(K{row}-H{row})*1000', f'=(K{row}/H{row}-1)*1000000',
            )
            for column, value in enumerate(cells, start=1):
                sheet.cell(row=row, column=column, value=value)
        last = row

        row += 2
        sheet.cell(row=row, column=1, value='Step 3: panel').font = _BOLD
        row += 1
        _header(sheet, row, ('', 'Panel X (um)', 'Panel X (ppm)', 'Panel Y (um)', 'Panel Y (ppm)',
                             'X - Y (ppm)', 'Quadrant range X (ppm)', 'Quadrant range Y (ppm)'))
        row += 1
        formulas = ('From the formulas',
                    f'=AVERAGE(F{first}:F{last})', f'=AVERAGE(G{first}:G{last})',
                    f'=AVERAGE(L{first}:L{last})', f'=AVERAGE(M{first}:M{last})',
                    f'=C{row}-E{row}',
                    f'=MAX(G{first}:G{last})-MIN(G{first}:G{last})',
                    f'=MAX(M{first}:M{last})-MIN(M{first}:M{last})')
        for column, value in enumerate(formulas, start=1):
            sheet.cell(row=row, column=column, value=value).fill = _RESULT
        row += 1
        panel, uniformity = group.result.full_panel.iloc[0], group.result.uniformity.iloc[0]
        app = ('Shown in the app',
               panel['Delta X (um)'], panel['Shrinkage X (ppm)'],
               panel['Delta Y (um)'], panel['Shrinkage Y (ppm)'], panel['X-Y Offset (ppm)'],
               uniformity['QTQ Range X (ppm)'], uniformity['QTQ Range Y (ppm)'])
        for column, value in enumerate(app, start=1):
            sheet.cell(row=row, column=column,
                       value=value if isinstance(value, str) else float(value))
        row += 3

    if not written:
        sheet.cell(row=row, column=1, value='No shrinkage result in scope, so there is nothing to calculate.')
    for column in range(2, len(_STEP_HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 20
        for cells in sheet.iter_rows(min_row=len(_EXPLAIN) + 2, min_col=column, max_col=column):
            cells[0].number_format = '0.000'
    sheet.column_dimensions['A'].width = 22


def write_shrinkage_excel(writer: pd.ExcelWriter, df: pd.DataFrame, settings: dict,
                          meta: dict | None = None, warn: WarnFn | None = None) -> None:
    """Write the registration shrinkage workbook for the rows in ``df`` (one file type)."""
    if df is None or df.empty:
        pd.DataFrame({'info': ['No data available for export scope.']}).to_excel(
            writer, sheet_name='Summary', index=False)
        return

    # Same guard as the main export: no text cell may become a formula.
    groups = compute_shrinkage_by_group(_sanitize_frame(df), settings)

    _write(writer, 'Summary', summary_frame(groups), 'No groups in scope.')
    _write(writer, 'Quadrants', _stack(groups, lambda g: g.result.quadrant),
           'No shrinkage result in scope.')
    _write(writer, 'Quadrant Sizes', _stack(groups, lambda g: g.result.lengths),
           'No quadrant sizes in scope.')

    _write_calculation(writer.book.create_sheet('Calculation'), groups)

    charts = writer.book.create_sheet('Charts')
    row = 1
    for group in groups:
        if group.result.quadrant.empty:
            continue
        charts.cell(row=row, column=1, value=_safe_text(f'{group.process} | {group.label}'))
        for column, (name, figure) in zip(('A', 'J'), (
                ('Panel map', plot_shrinkage_panel(group.result, group.design, 'Shrinkage map',
                                                   height=_CHART_HEIGHT)),
                ('Quadrant bars', plot_quadrant_bars(group.result, 'Shrinkage per quadrant',
                                                     height=_CHART_HEIGHT)))):
            anchor = f'{column}{row + 1}'
            try:
                charts.add_image(_fig_to_xl_image(figure, _CHART_WIDTH, _CHART_HEIGHT), anchor)
            except Exception as exc:  # no Chrome / kaleido: keep the workbook, mark the chart
                logger.warning('Shrinkage chart "%s" could not be rendered: %s', name, exc)
                _notify(warn, f'Shrinkage chart "{name}" could not be rendered: {exc}')
                _write_chart_fallback(charts, anchor, name, exc)
        row += _ROWS_PER_CHART
    if row == 1:
        charts.cell(row=1, column=1, value='No shrinkage result in scope, so there are no charts.')

    notes = writer.book.create_sheet('Notes')
    lines = [
        'Registration shrinkage',
        f'Method: {METHOD_LABEL}, as in the company workbook: per quadrant the two X and two Y '
        f'lengths against design, panel = mean of the four quadrants.',
        'Sign: measured minus design, so negative = shrink.',
        'um = change over one quadrant length; ppm = that change divided by the design length.',
        f"Outlier alert: a quadrant more than {float(shrinkage_config(settings).get('outlier_ppm', 50)):.0f} "
        f"ppm from the panel average.",
    ]
    for key in ('project_name', 'part_number', 'lot_number', 'lot_name'):
        if (meta or {}).get(key):
            lines.append(f"{key.replace('_', ' ').title()}: {meta[key]}")
    for group in groups:
        lines.append('')
        lines.append(f'{group.process} | {group.label} | {group.point} centres')
        if group.lengths:
            lines.append(f'Design quadrant size: {group.lengths[0]:g} x {group.lengths[1]:g} mm')
        lines.extend(group.result.warnings)
    for index, line in enumerate(lines, start=1):
        notes.cell(row=index, column=1, value=_safe_text(line))
