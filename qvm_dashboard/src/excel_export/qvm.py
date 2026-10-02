"""QVM (Pad-to-Via / Via-to-Pad) workbook writer."""

import logging

import pandas as pd

from src.display_frame import GRID_ID_HEADER, build_display_frame
from src.registration_shrinkage import compute_shrinkage_by_group, shrinkage_config
from src.utils import is_via_to_pad
from src.visuals import (
    plot_vtp_bullseye, plot_vtp_site_bars, plot_vtp_box, plot_quiver,
)
from ._shared import (
    WarnFn,
    _REGION_MAP,
    _add_panel_map,
    _fig_to_xl_image,
    _files_sheet,
    _notify,
    _safe_sheet_name,
    _safe_text,
    _sanitize_frame,
    _warnings_sheet,
    _write_chart_fallback,
)

logger = logging.getLogger(__name__)


def _add_analytics_sheet(writer: pd.ExcelWriter, df: pd.DataFrame,
                         settings: dict,
                         sheet_label: str = 'Analytics',
                         warn: WarnFn | None = None) -> None:
    """Create an analytics sheet with one horizontal 4-chart row per side.

    Layout:
        A1   — title label (includes sheet_label for self-identification)
        For each available side (Front then Back), place one horizontal row:
            <start>3    — Bullseye scatter (DX/DY)
            <start+13>3 — Per-site SC bars
            <start+26>3 — Box plot (SC distribution)
            <start+39>3 — Quiver (ABF distortion map)

    sheet_label is used as the worksheet name so that when both F and B sides
    are uploaded for the same Build Up, one sheet can still show both sides as
    separate horizontal rows, avoiding mixed/combined charts.

    Each chart is embedded independently so one failed render (or missing
    kaleido) never aborts the rest of the sheet or the workbook.
    """
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        logger.warning('Analytics sheet skipped — openpyxl missing: %s', exc)
        _notify(warn, '⚠️ Analytics sheet not embedded: `openpyxl` required.')
        return

    ws = writer.book.create_sheet(sheet_label)
    ws['A1'] = f'Registration Analytics — {sheet_label}'
    ws['A1'].font = openpyxl.styles.Font(bold=True, size=12)

    multiplier = settings.get('PLOT_SETTINGS', {}).get('vector_multiplier_default', 100)

    side_order = {'F': 0, 'B': 1}
    side_labels = {'F': 'Front', 'B': 'Back'}
    side_values = []
    if 'Side' in df.columns:
        side_values = [
            str(v).strip().upper() for v in df['Side'].dropna().tolist()
            if str(v).strip()
        ]
        side_values = sorted(set(side_values), key=lambda s: (side_order.get(s, 2), s))
    if not side_values:
        side_values = ['ALL']

    # One side row spans four charts across the sheet.
    # Start each additional side row further down so Front/Back are stacked.
    from openpyxl.utils.cell import get_column_letter

    def _anchor(start_col: int, row: int) -> str:
        return f'{get_column_letter(start_col)}{row}'

    for block_idx, side in enumerate(side_values):
        start_row = 3 + (block_idx * 72)
        start_col = 2
        side_df = df if side == 'ALL' else df[df['Side'].astype(str).str.upper() == side].copy()
        if side_df.empty:
            continue

        side_label = side_labels.get(side, side)
        ws[_anchor(1, start_row - 1)] = f'Side: {side_label}'
        ws[_anchor(1, start_row - 1)].font = openpyxl.styles.Font(bold=True)

        charts = [
            (_anchor(start_col,      start_row), 'Bullseye',  lambda d=side_df: plot_vtp_bullseye(d, settings),    500, 500),
            (_anchor(start_col + 13, start_row), 'Site bars', lambda d=side_df: plot_vtp_site_bars(d, settings),   700, 500),
            (_anchor(start_col + 26, start_row), 'Box plot',  lambda d=side_df: plot_vtp_box(d, settings),         700, 500),
            (_anchor(start_col + 39, start_row), 'Quiver',    lambda d=side_df: plot_quiver(d, settings, multiplier), 500, 500),
        ]
        for anchor, name, make_fig, w, h in charts:
            chart_name = f'{name} [{side}]'
            try:
                ws.add_image(_fig_to_xl_image(make_fig(), width_px=w, height_px=h), anchor)
            except Exception as exc:
                logger.warning('%s chart embed failed for %s: %s',
                               chart_name, sheet_label, exc, exc_info=True)
                _write_chart_fallback(ws, anchor, chart_name, exc)
                _notify(warn,
                        f'⚠️ {chart_name} not embedded in "{sheet_label}": {exc}')


def _process_slices(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Split scoped data into VTP/PTV dataframes for unified exports."""
    vtp_mask = is_via_to_pad(df)
    return {
        'VTP': df[vtp_mask].copy(),
        'PTV': df[~vtp_mask].copy(),
    }


def _iter_build_up_frames(proc_df: pd.DataFrame) -> list[tuple[str, pd.DataFrame]]:
    if proc_df.empty:
        return []
    if 'Build Up' in proc_df.columns:
        bu_values = sorted({str(v).strip() for v in proc_df['Build Up'].dropna().tolist() if str(v).strip()})
    else:
        bu_values = []
    if not bu_values:
        bu_values = ['Unknown']

    out = []
    for bu in bu_values:
        if 'Build Up' in proc_df.columns:
            if bu.lower() == 'unknown':
                bu_df = proc_df[proc_df['Build Up'].astype(str).str.strip().str.lower() == 'unknown'].copy()
            else:
                bu_df = proc_df[proc_df['Build Up'] == bu].copy()
        else:
            bu_df = proc_df.copy()
        if not bu_df.empty:
            out.append((bu, bu_df))
    return out


def _write_shrinkage_summary_sheet(
    writer: pd.ExcelWriter,
    settings: dict,
    process_frames: dict[str, pd.DataFrame],
    used_sheet_names: set[str],
) -> bool:
    """One sheet with the registration shrinkage of every Build Up / Panel / Side.

    Uses the report method (``REGISTRATION_SHRINKAGE.default_method``, normally
    the company calculation) and the design positions configured for each
    Build Up and file type. A group without design positions is listed with
    the reason instead of numbers.
    """
    sections = []
    for process_code, proc_df in process_frames.items():
        if proc_df.empty:
            continue
        for group in compute_shrinkage_by_group(proc_df, settings):
            sections.append((process_code, group))
    if not sections:
        return False

    summary_rows = []
    for process_code, group in sections:
        result = group.result
        if result.full_panel.empty:
            continue
        panel = result.full_panel.iloc[0]
        uniformity = result.uniformity.iloc[0]
        row = {'Process': process_code, **group.group}
        row.update({
            'FP X (um)': panel['Delta X (um)'],
            'FP Y (um)': panel['Delta Y (um)'],
            'FP X (ppm)': panel['Shrinkage X (ppm)'],
            'FP Y (ppm)': panel['Shrinkage Y (ppm)'],
            'X-Y Offset (ppm)': panel['X-Y Offset (ppm)'],
            'QTQ X (ppm)': uniformity['QTQ Range X (ppm)'],
            'QTQ Y (ppm)': uniformity['QTQ Range Y (ppm)'],
            'Alert Count': len(result.outliers),
        })
        if 'Status' in panel:
            row['Status'] = panel['Status']
        summary_rows.append(row)

    method_label = sections[0][1].result.method_label
    outlier_ppm = float(shrinkage_config(settings).get('outlier_ppm', 50.0))
    sheet_name = _safe_sheet_name('Shrinkage Summary', used_sheet_names)
    ws = writer.book.create_sheet(sheet_name)
    ws['A1'] = 'Registration Shrinkage Summary (PTV + VTP)'
    ws['A2'] = (f'Method: {method_label}. Measured minus design (negative = shrink). '
                f'um = change over one quadrant length.')

    row = 4
    if summary_rows:
        summary_df = pd.DataFrame(summary_rows).round(3)
        summary_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row - 1)
        row += len(summary_df) + 2

    for process_code, group in sections:
        result = group.result
        ws.cell(row=row, column=1, value=f'{process_code} | {group.label}')
        row += 1
        ws.cell(row=row, column=1,
                value=f'{group.process}, {group.point} centres, {result.method_label}')
        row += 1
        if group.design_note:
            ws.cell(row=row, column=1, value=group.design_note)
            row += 1

        if not result.quadrant.empty:
            ws.cell(row=row, column=1, value='Full Panel')
            row += 1
            result.full_panel.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row - 1)
            row += len(result.full_panel) + 2

            ws.cell(row=row, column=1, value='Quadrants')
            row += 1
            result.quadrant.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row - 1)
            row += len(result.quadrant) + 2

            ws.cell(row=row, column=1, value='Uniformity')
            row += 1
            result.uniformity.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row - 1)
            row += len(result.uniformity) + 2

            ws.cell(row=row, column=1, value='Outlier Alerts')
            row += 1
            if result.outliers.empty:
                ws.cell(row=row, column=1,
                        value=f'No quadrant is more than {outlier_ppm:.0f} ppm from the panel average.')
                row += 2
            else:
                result.outliers.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row - 1)
                row += len(result.outliers) + 2

        if result.quadrant.empty and not result.lengths.empty:
            ws.cell(row=row, column=1, value='Measured quadrant sizes (no shrinkage reported)')
            row += 1
            result.lengths.to_excel(writer, sheet_name=sheet_name, index=False, startrow=row - 1)
            row += len(result.lengths) + 2

        if result.warnings:
            ws.cell(row=row, column=1, value='Warnings')
            row += 1
            for msg in result.warnings:
                ws.cell(row=row, column=1, value=msg)
                row += 1
        row += 1

    return True


def _report_sheet(df: pd.DataFrame, grid_id_col: str, meta: dict,
                  usl_col_name: str, usl_value: float,
                  include_xy_usl: bool = False) -> pd.DataFrame:
    """Full report frame: metadata prepended, USL reference column(s) appended.

    Per-row provenance (Build Up, Panel, Side) is read from ``df`` itself, which
    build_display_frame has already sorted, so each value stays with its own
    fiducial. Taking them from a separately-ordered source misattributed them:
    the display frame is sorted by Grid ID while the parsed frame is grouped by
    file, so a scope covering both sides interleaved the two orders and half the
    rows were labelled with the other side's values.

    Job-level metadata (project, part, lot) has no per-row variation and comes
    from ``meta``. `usl_col_name`/`usl_value` name the process-specific
    spec-limit column ('VtP_USL' for Via-to-Pad, 'PtV_USL' for Pad-to-Via). When
    `include_xy_usl` is set, blank X/Y USL columns are appended too
    (Via-to-Pad only).
    """
    s = df.copy()

    def _per_row(column: str, meta_key: str):
        """That column's own values when present, else the job-level fallback."""
        return s[column] if column in s.columns else meta.get(meta_key, '')

    side_vals     = _per_row('Side', 'side')
    panel_vals    = _per_row('Panel', 'panel')
    build_up_vals = _per_row('Build Up', 'build_up')

    drop_if_present = [
        'Project_Name', 'Part_Number', 'Lot_Number', 'Lot_Name',
        'Buildup', 'Build Up', 'Process', 'Panel', 'Side',
        'Fiducial', 'Fiducial_Number', 'VtP_USL', 'PtV_USL', 'X_USL', 'Y_USL',
    ]
    existing = [c for c in drop_if_present if c in s.columns]
    if existing:
        s = s.drop(columns=existing)

    # Keep report generation resilient when Grid ID is missing/malformed.
    if grid_id_col in s.columns:
        grid_str = s[grid_id_col].astype(str).str.replace(r'\D', '', regex=True).str.zfill(2)
        s.insert(0, 'Fiducial', grid_str.str[0].map(_REGION_MAP))
        s.insert(0, 'Fiducial_Number', pd.to_numeric(grid_str.str[1], errors='coerce'))
    else:
        logger.warning('Report export: missing grid id column "%s"; fiducial fields left blank', grid_id_col)
        s.insert(0, 'Fiducial', None)
        s.insert(0, 'Fiducial_Number', None)

    # Job-level metadata comes from sidebar text inputs: neutralise formula
    # injection at insertion (the frame was sanitised before this ran).
    s.insert(0, 'Side',        side_vals)
    s.insert(0, 'Panel',       panel_vals)
    s.insert(0, 'Process',     _safe_text(meta.get('process', '')))
    s.insert(0, 'Buildup',     build_up_vals)
    s.insert(0, 'Lot_Name',    _safe_text(meta.get('lot_name', '')))
    s.insert(0, 'Lot_Number',  _safe_text(meta.get('lot_number', '')))
    s.insert(0, 'Part_Number', _safe_text(meta.get('part_number', '')))
    s.insert(0, 'Project_Name', _safe_text(meta.get('project_name', '')))
    if grid_id_col in s.columns:
        s = s.drop(columns=[grid_id_col])

    # Append USL reference column(s) (X/Y USL intentionally left blank)
    s[usl_col_name] = usl_value
    if include_xy_usl:
        s['X_USL'] = None
        s['Y_USL'] = None

    return s


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def write_unified_excel(writer: pd.ExcelWriter, scoped_df: pd.DataFrame,
                        settings: dict, meta: dict,
                        warnings: list = None,
                        warn: WarnFn | None = None) -> None:
    """Write one workbook with process-separated sheets and BU analytics.

    Sheets written:
    - 'Via to Pad' — process data with panel map
    - 'Pad to Via' — process data with panel map
    - 'Report' — consolidated process-aware report table
    - 'Analytics <PROC> <BU>' — one analytics sheet per process and Build Up
    - 'Shrinkage Summary' — one combined shrinkage sheet for all process/BU sections
    - 'Warnings' — parser warnings
    - 'Files' — input-file manifest (only when ``meta['file_records']`` is set)

    ``warn`` is the user-facing warning sink (the app passes
    ``st.sidebar.warning``); ``None`` means log-only.
    """
    if scoped_df is None or scoped_df.empty:
        pd.DataFrame({'info': ['No data available for export scope.']}).to_excel(
            writer, sheet_name='Report', index=False)
        ws_warn = writer.book.create_sheet('Warnings')
        _warnings_sheet(ws_warn, warnings or [])
        return

    # Neutralise formula injection at the write layer: every sheet below
    # derives from this frame, so one sanitised copy covers process sheets,
    # Report, and analytics tables. In-memory analysis frames are untouched.
    scoped_df = _sanitize_frame(scoped_df)

    process_frames = _process_slices(scoped_df)
    vtp_usl = float(settings.get('VTP_REGISTRATION', {}).get('vtp_usl_um', 10.0))
    ptv_usl = float(settings.get('PTV_REGISTRATION', {}).get('ptv_usl_um', 10.0))

    report_frames = []
    for process_code, proc_df in process_frames.items():
        if proc_df.empty:
            continue

        is_vtp = process_code == 'VTP'
        sheet_base = 'Via to Pad' if is_vtp else 'Pad to Via'
        proc_display = build_display_frame(proc_df, settings, is_via_to_pad=is_vtp)
        proc_meta = dict(meta)
        proc_meta['process'] = sheet_base

        proc_display.to_excel(writer, sheet_name=sheet_base, index=False)
        _add_panel_map(writer.sheets[sheet_base], settings, sheet_base, warn)

        report_frames.append(
            _report_sheet(
                proc_display,
                GRID_ID_HEADER,
                proc_meta,
                usl_col_name='VtP_USL' if is_vtp else 'PtV_USL',
                usl_value=vtp_usl if is_vtp else ptv_usl,
                include_xy_usl=is_vtp,
            )
        )

    if report_frames:
        pd.concat(report_frames, ignore_index=True).to_excel(
            writer, sheet_name='Report', index=False)
    else:
        pd.DataFrame({'info': ['No process-specific rows found.']}).to_excel(
            writer, sheet_name='Report', index=False)

    used_sheet_names = set(writer.book.sheetnames)
    for process_code, proc_df in process_frames.items():
        if proc_df.empty:
            continue

        for bu, bu_df in _iter_build_up_frames(proc_df):

            label = _safe_sheet_name(f'Analytics {process_code} {bu}', used_sheet_names)
            _add_analytics_sheet(writer, bu_df, settings, sheet_label=label, warn=warn)

    _write_shrinkage_summary_sheet(writer, settings, process_frames, used_sheet_names)

    ws_warn = writer.book.create_sheet(_safe_sheet_name('Warnings', used_sheet_names))
    _warnings_sheet(ws_warn, warnings or [])

    file_records = (meta or {}).get('file_records')
    if file_records is not None:
        ws_files = writer.book.create_sheet(
            _safe_sheet_name('Files', used_sheet_names))
        _files_sheet(ws_files, file_records)
