"""Shared Excel plumbing: Chrome discovery, PNG rendering, panel map, warnings.

Streamlit-free by design — user-facing warnings go through the ``warn``
callback injected by the caller (the app passes ``st.sidebar.warning``),
so this package imports in tests and scripts without a Streamlit runtime.
"""

import logging
import os
import re
from collections.abc import Callable
from io import BytesIO

import pandas as pd
import plotly.graph_objects as go
from openpyxl.drawing.image import Image as XLImage

logger = logging.getLogger(__name__)

# Quarter prefix → label for the Report sheet Fiducial column
_REGION_MAP = {'1': 'UL', '2': 'LL', '3': 'LR', '4': 'UR'}

#: User-facing warning sink, e.g. ``st.sidebar.warning``. ``None`` means
#: log-only (every site already logs via ``logger``).
WarnFn = Callable[[str], None]

#: Absolute path of the ``qvm_dashboard/`` package root, so asset paths from
#: settings (e.g. ``EXCEL_EXPORT.panel_map_image``) resolve identically for
#: the app, the tests, and this package's new nesting depth.
PACKAGE_ROOT = os.path.dirname(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def _notify(warn: WarnFn | None, message: str) -> None:
    if warn is not None:
        warn(message)


def _ensure_chrome_on_path() -> None:
    """Make Chrome discoverable to Kaleido when it is installed but not on PATH.

    This is most common on Windows, where Chrome can exist under one of the
    standard Program Files locations without being on the system PATH. Adding the
    chrome installation directory to PATH and exporting GOOGLE_CHROME_BIN helps
    Plotly/Kaleido locate the browser during fig.to_image().
    """
    candidates = [
        r'C:\Program Files\Google\Chrome\Application\chrome.exe',
        r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe',
        os.path.expanduser(r'~\AppData\Local\Google\Chrome\Application\chrome.exe'),
        os.path.expanduser(r'~\AppData\Local\Google\Chrome SxS\Application\chrome.exe'),
        r'C:\Program Files\Chromium\Application\chrome.exe',
        r'C:\Program Files (x86)\Chromium\Application\chrome.exe',
    ]
    for candidate in candidates:
        if os.path.isfile(candidate):
            # Separator-agnostic dirname: the candidates are Windows paths,
            # and os.path.dirname misses them entirely on POSIX hosts.
            chrome_dir = re.sub(r'[\\/][^\\/]+$', '', candidate)
            path_entries = os.environ.get('PATH', '').split(os.pathsep)
            if chrome_dir not in path_entries:
                os.environ['PATH'] = chrome_dir + os.pathsep + os.environ.get('PATH', '')
            os.environ['GOOGLE_CHROME_BIN'] = candidate
            os.environ['CHROME_BIN'] = candidate
            return


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _fig_to_xl_image(fig: go.Figure, width_px: int = 300,
                     height_px: int = 400) -> XLImage:
    """Render a Plotly figure to PNG via kaleido and wrap it as an XLImage."""
    _ensure_chrome_on_path()
    img_bytes = fig.to_image(format='png', width=width_px, height=height_px, scale=2)
    buf = BytesIO(img_bytes)
    xl = XLImage(buf)
    xl.width = width_px
    xl.height = height_px
    return xl


def _write_chart_fallback(ws, anchor: str, chart_name: str, exc: Exception) -> None:
    """Embed a lightweight placeholder PNG when a chart image cannot be rendered."""
    try:
        from PIL import Image as PILImage, ImageDraw, ImageFont

        width, height = 400, 220
        image = PILImage.new('RGB', (width, height), color=(255, 247, 247))
        draw = ImageDraw.Draw(image)
        title = f'{chart_name} unavailable'
        message = 'Chart rendering failed; workbook preserved without live graphic.'
        detail = str(exc)

        try:
            font = ImageFont.load_default()
        except Exception:
            font = None

        draw.text((20, 20), title, fill=(156, 0, 6), font=font)
        draw.text((20, 55), message, fill=(90, 90, 90), font=font)
        draw.text((20, 85), detail[:110], fill=(120, 120, 120), font=font)

        buf = BytesIO()
        image.save(buf, format='PNG')
        buf.seek(0)
        xl = XLImage(buf)
        xl.width = width
        xl.height = height
        ws.add_image(xl, anchor)
    except Exception:
        cell = ws[anchor]
        cell.value = f"{chart_name} chart unavailable: {exc}"
        try:
            import openpyxl.styles as xl_styles
            cell.font = xl_styles.Font(italic=True, color='9C0006')
            cell.fill = xl_styles.PatternFill(fill_type='solid', fgColor='FFC7CE')
        except ImportError:
            pass


def _load_panel_map_image(settings: dict) -> XLImage | None:
    """Load Panel_Map.png from the path configured in EXCEL_EXPORT settings.

    Returns None (with a logged warning) if the file does not exist.
    """
    export_cfg = settings.get('EXCEL_EXPORT', {})
    rel_path   = export_cfg.get('panel_map_image', 'assets/Panel_Map.png')
    panel_cm   = float(export_cfg.get('panel_image_cm', 8.5))
    panel_px   = int(panel_cm / 2.54 * 96)  # cm → px at 96 dpi

    # Resolve relative to the qvm_dashboard package root.
    abs_path = os.path.join(PACKAGE_ROOT, rel_path)

    if not os.path.isfile(abs_path):
        logger.warning('Panel map image not found: %s', abs_path)
        return None

    xl = XLImage(abs_path)
    xl.width  = panel_px
    xl.height = panel_px
    return xl


# ---------------------------------------------------------------------------
# Sheet-level chart embedders
# ---------------------------------------------------------------------------

def _add_panel_map(ws, settings: dict, label: str, warn: WarnFn | None = None) -> None:
    """Embed Panel_Map.png to the right of the data table to avoid overlap.

    Measurement sheets carry only the data table and this reference image; the
    bullseye, SC bars, box, and quiver charts live solely on the Analytics
    sheets built by :func:`src.excel_export.qvm._add_analytics_sheet`.
    """
    try:
        from openpyxl.utils.cell import get_column_letter

        xl = _load_panel_map_image(settings)
        if xl is None:
            _notify(warn, f"⚠️ Panel_Map.png not found — skipped in '{label}'.")
            return

        data_end_col = max(int(getattr(ws, 'max_column', 1) or 1), 1)
        anchor_col = get_column_letter(data_end_col + 2)
        ws.add_image(xl, f'{anchor_col}1')
    except Exception as exc:
        logger.warning('Panel map embed failed for %s: %s', label, exc, exc_info=True)
        _notify(warn, f"⚠️ Panel map not embedded in '{label}': {exc}")


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _warnings_sheet(ws, warnings: list) -> None:
    """Write parser warnings into a worksheet, one per row.

    Header row is bold with a red background. If there are no warnings, a
    placeholder 'No warnings for this file.' message is written.
    """
    try:
        import openpyxl.styles as xl_styles
    except ImportError:
        return

    header_fill = xl_styles.PatternFill(
        fill_type='solid', fgColor='C00000')  # dark red
    header_font = xl_styles.Font(bold=True, color='FFFFFF')

    ws.column_dimensions['A'].width = 100
    ws['A1'] = 'Warnings'
    ws['A1'].fill = header_fill
    ws['A1'].font = header_font

    if not warnings:
        ws['A2'] = 'No warnings for this file.'
    else:
        for i, msg in enumerate(warnings, start=2):
            ws.cell(row=i, column=1, value=msg)


def _safe_sheet_name(sheet_name: str, existing: set[str]) -> str:
    """Return an Excel-safe, unique worksheet name (<=31 chars)."""
    base = re.sub(r'[\[\]:*?/\\]', '-', str(sheet_name).strip()) or 'Sheet'
    base = re.sub(r'\s+', ' ', base)
    candidate = base[:31]
    counter = 2
    while candidate in existing:
        suffix = f' {counter}'
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    existing.add(candidate)
    return candidate
